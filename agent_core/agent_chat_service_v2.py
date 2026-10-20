"""
Agent Chat Service V2 - Agent聊天服务核心（状态机版本）

使用状态机模式重构的 Agent 聊天服务：
- while(true) + Continue/Terminal 控制流
- 工具智能并发执行
- 完整的流式输出
- 状态跨迭代传递

核心改进：
1. 状态机模式替代简单 for 循环
2. 只读工具并行执行，写工具串行执行
3. 更清晰的错误恢复机制

完整功能（从原版迁移）：
- 上下文压缩（自动/响应式）
- 记忆系统（Auto Memory / Agent Memory）
- 标题自动生成
- Transcript 管理
- Todo 回调
- Ask Question 阻塞机制
- 权限检查
- 文件/技能追踪
"""

import asyncio
import logging
import json
import uuid
import os
from datetime import datetime, timezone, timedelta
from typing import List, Dict, Any, Optional, AsyncGenerator
from openai import AsyncOpenAI


from .agent_loader import get_cached_agent, filter_tools_for_agent, list_available_agents_from_yaml
from .llm_client import apply_reasoning_mode, build_default_headers, keep_max_effort, shaping_from_provider
from .context_compressor import (
    # Token 计数（tiktoken + 中英文比例降级）
    estimate_tokens,
)
# 导入 Transcript 管理器
from .transcript_manager import (
    get_transcript_manager,
)
from . import agent_chat_history as history

# 导入经验账本
from .memory import get_ledger

from .tool_definitions import get_tools_definition, get_tools_definition_with_mcp
from .tool_executor_impl import AgentToolExecutorImpl

# MCP 管理器（延迟导入，避免循环依赖）
_mcp_manager = None

from .chat.usage import UsageTracker
from .chat.stream_client import StreamConsumer, open_stream
from .chat.tool_cycle import ToolCycle
from .chat.turn_executor import TurnState, prepare_turn

logger = logging.getLogger(__name__)

# ==================== 模型上下文限制 ====================

MODEL_CONTEXT_LIMITS = {
    'deepseek': 1_000_000,
    'mimo': 1_000_000,
    'minimax': 1_000_000,
    'gpt-4': 128_000,
    'gpt-3.5': 16_000,
    'sonnet': 200_000,
    'opus': 200_000,
    'haiku': 200_000,
    'gemini': 1_000_000,
}

def get_model_context_limit(model_name: str) -> int:
    """根据模型名称获取上下文限制 token 数"""
    name = model_name.lower().replace(' ', '')
    for key, limit in MODEL_CONTEXT_LIMITS.items():
        if key in name:
            return limit
    return 128_000  # default fallback（保守窗口，避免未知模型高估）

# ==================== Token 计数辅助 ====================

def count_tokens(text: str) -> int:
    """计算 token 数（统一委托 context_compressor.estimate_tokens）"""
    return estimate_tokens(text)

def count_tokens_for_full_context(
    system_prompt: str,
    messages: list,
    tools_definitions: list
) -> int:
    """计算完整上下文（system prompt + 历史消息 + tools 定义）的总 token 数"""
    total = count_tokens(system_prompt or '')
    for msg in messages:
        content = msg.get('content') or ''
        total += count_tokens(content)
    # tools 定义为 JSON 序列化后计数
    if tools_definitions:
        import json
        tools_json = json.dumps(tools_definitions, ensure_ascii=False)
        total += count_tokens(tools_json)
    return total


def _is_suspicious_stream_end(finish_reason: Optional[str], response_text: str) -> bool:
    """可疑流结束检测。

    在"无工具调用、即将结束回合"时判定，两类可疑特征：
    1. finish_reason == 'length'：输出到达 max_tokens 被截断，任务未必完成；
    2. 无文字输出：模型在"思考→行动"衔接点静默停止（特征：reasoning 正常
       生成、本应发出工具调用或文字收尾，却直接断流）、完全空响应（API 异常）、
       或工具调用无效被全部丢弃等场景。
    正常结束（finish_reason=stop 且有文字收尾）不在此列。
    """
    if finish_reason == "length":
        return True
    if not (response_text or "").strip():
        return True
    return False

# ==================== Agent Chat Service V2 ====================

_DATA_HOME = os.environ.get("PEBBLE_HOME") or os.path.join(os.path.expanduser("~"), ".pebble")

# ==================== 高峰/空闲时段判定（北京时间） ====================
# DeepSeek 分时定价规则（前端 StatusBar 展示与后端计费必须使用同一规则）：
#   高峰时段 = 北京时间周一至周五（不含中国法定节假日）09:00-12:00、14:00-18:00
#   （左闭右开：[09:00,12:00) ∪ [14:00,18:00)，12:00/18:00 整点转入空闲）
#   空闲时段 = 其余全部时间（含周末、中国法定节假日全天、12:00-14:00、18:00 后、09:00 前）
# 中国法定节假日 = 国务院公告的放假安排（含调休连休），数据在 cn_holidays.json：
#   放假日期命中即全天空闲；调休补班日按官方文字口径不参与判定（周末补班日仍按周末全天空闲处理）。
# 空闲价格 = 高峰价格的一半（由 model_pricing.json 的 pricing_offpeak 显式配置）。
# 北京时间 = UTC+8 固定偏移（无夏令时），不依赖服务器本地时区。

def is_peak_period(now: Optional[datetime] = None) -> bool:
    """判断给定时刻是否处于高峰计费时段（默认当前时刻，北京时间）"""
    if now is None:
        now = datetime.now()
    bj = now.astimezone(timezone(timedelta(hours=8)))
    if bj.weekday() >= 5:  # 周六(5)/周日(6)
        return False
    if is_cn_holiday(bj):  # 中国法定节假日（国务院公告放假安排，含调休连休）全天空闲
        return False
    h = bj.hour
    return (9 <= h < 12) or (14 <= h < 18)


def current_pricing_period(now: Optional[datetime] = None) -> str:
    """返回当前计费时段标识：'peak'（高峰）或 'offpeak'（空闲）"""
    return "peak" if is_peak_period(now) else "offpeak"


def _read_json_config(filename: str) -> Optional[dict]:
    """读取 JSON 配置，优先用户数据目录，回退到打包根"""
    import os as _os
    data_path = _os.path.join(_DATA_HOME, filename)
    bundled_path = _os.path.join(_os.path.dirname(_os.path.dirname(__file__)), filename)
    for p in [data_path, bundled_path]:
        if _os.path.exists(p):
            try:
                with open(p, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                continue
    return None


def build_thinking_params(model: str, thinking_enabled: bool, reasoning_effort: str, keep_max: bool = False) -> Dict[str, Any]:
    """构建思考模式相关的请求参数（所有厂商统一处理）。

    - 开启：thinking={"type": "enabled"} + reasoning_effort（全厂商发送，不再限 DeepSeek）
      - 非 DeepSeek 模型的 "max" 档降级为 "high"：小米 MiMo 仅接受 low/medium/high
        （max 会 400），OpenAI/Gemini 值域亦不含 max；provider 配 maxEffort=keep
        时跳过降级（如 OpenCode Go 的 GLM-5.3 支持 max）
    - 关闭：thinking={"type": "disabled"}（显式关闭，不再发送 reasoning_effort）
    """
    if not thinking_enabled:
        return {"extra_body": {"thinking": {"type": "disabled"}}}
    effort = reasoning_effort
    if not keep_max and effort == "max" and "deepseek" not in (model or "").lower():
        effort = "high"
    return {
        "extra_body": {"thinking": {"type": "enabled"}},
        "reasoning_effort": effort,
    }

# ==================== 中国法定节假日数据（cn_holidays.json） ====================
# 用途：is_peak_period 判定"周一至周五（不含中国法定节假日）"中的节假日部分。
# 数据：国务院公告放假安排（含调休连休）——holidays=放假日列表（命中即全天空闲），
#   workdays=调休补班日列表（按官方文字口径不参与判定，仅记录）。
# 加载：路径优先级 ~/.pebble/cn_holidays.json > 打包根/cn_holidays.json；候选文件快照比对热重载（免重启，同 model_pricing 理念）。
# 降级：数据缺失/损坏时按空集合处理（判定退化为周末规则），保证系统可用性。

_cn_holidays: set = set()          # 放假日日期集合（"YYYY-MM-DD"）
_cn_workdays: set = set()          # 补班日日期集合（记录用，不参与判定）
_cn_holidays_meta: dict = {}       # 数据文件元信息（years_covered / updated_at 等）
_cn_holidays_watch: dict = {}      # 候选文件状态快照（path -> mtime），供热重载比对
_cn_holidays_warned: bool = False  # 数据缺失/损坏警告是否已提示（防刷屏）

def _cn_holidays_candidates():
    """节假日数据文件候选路径（与 _read_json_config 相同的优先级：数据目录 > 打包根）"""
    return [
        os.path.join(_DATA_HOME, "cn_holidays.json"),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), "cn_holidays.json"),
    ]

def _snapshot_cn_holiday_files() -> dict:
    """候选数据文件状态快照（path -> mtime），供热重载比对"""
    snap = {}
    for p in _cn_holidays_candidates():
        try:
            if os.path.exists(p):
                snap[p] = os.path.getmtime(p)
        except OSError as e:
            logger.debug("cn_holidays file stat failed (%s): %s", p, e)
    return snap

def _load_cn_holidays() -> None:
    """读取 cn_holidays.json 并重建模块级缓存；失败时降级为空集合。
    解析失败自动尝试下一个候选（与 _read_json_config 同模式）。"""
    global _cn_holidays, _cn_workdays, _cn_holidays_meta, _cn_holidays_warned, _cn_holidays_watch
    data = None
    for p in _cn_holidays_candidates():
        if os.path.exists(p):
            try:
                with open(p, "r", encoding="utf-8") as f:
                    data = json.load(f)
                break
            except Exception as e:
                logger.warning("cn_holidays.json 解析失败（%s）：%s", p, e)
                continue
    if isinstance(data, dict):
        holidays = data.get("holidays")
        workdays = data.get("workdays")
        _cn_holidays = {d for d in holidays if isinstance(d, str)} if isinstance(holidays, list) else set()
        _cn_workdays = {d for d in workdays if isinstance(d, str)} if isinstance(workdays, list) else set()
        _cn_holidays_meta = data
        _cn_holidays_warned = False
        if not _cn_holidays:
            logger.warning("cn_holidays.json 已加载但 holidays 为空，节假日判定不生效")
        # 年份覆盖检查：当前年份（北京时间）无数据时提醒更新（次年安排通常在 11 月发布）
        try:
            cur_year = datetime.now(timezone(timedelta(hours=8))).year
            covered = data.get("years_covered")
            if isinstance(covered, list) and cur_year not in covered:
                logger.warning("cn_holidays.json 未覆盖 %s 年，节假日判定本年不生效，请更新数据文件", cur_year)
        except Exception as e:
            logger.debug("cn_holidays year coverage check failed: %s", e)
    else:
        # 降级为空集合（周末规则）：数据缺失/损坏时判定退化为现状行为，保证系统可用
        _cn_holidays, _cn_workdays, _cn_holidays_meta = set(), set(), {}
        if not _cn_holidays_warned:
            _cn_holidays_warned = True
            logger.warning("cn_holidays.json 不可用（缺失或损坏），节假日判定降级为周末规则（中国法定节假日将按高峰计费）")
    _cn_holidays_watch = _snapshot_cn_holiday_files()

def _ensure_cn_holidays_fresh() -> None:
    """轻量热重载：任一候选文件出现/消失/mtime 变化时自动重读（模块级实现，server.py 直接 import 场景同样生效）。
    变更检测覆盖：修改文件、后补文件、删除文件、修复损坏文件；比对失败静默跳过（不影响判定）。"""
    try:
        if _snapshot_cn_holiday_files() != _cn_holidays_watch:
            _load_cn_holidays()
    except Exception as e:
        logger.debug("cn_holidays hot reload check failed: %s", e)

def is_cn_holiday(now: Optional[datetime] = None) -> bool:
    """判断（北京时间）给定时刻所在日期是否为中国法定节假日（放假日）。
    数据缺失/损坏时返回 False（降级为周末规则）。"""
    if now is None:
        now = datetime.now()
    bj = now.astimezone(timezone(timedelta(hours=8)))
    _ensure_cn_holidays_fresh()
    return bj.strftime("%Y-%m-%d") in _cn_holidays

def get_cn_holidays_payload() -> dict:
    """返回节假日数据负载（server.py 推送给前端做本地预判/展示用）"""
    _ensure_cn_holidays_fresh()
    years = _cn_holidays_meta.get("years_covered")
    return {
        "holidays": sorted(_cn_holidays),
        "workdays": sorted(_cn_workdays),
        "years_covered": years if isinstance(years, list) else [],
        "updated_at": _cn_holidays_meta.get("updated_at"),
    }

# 模块加载时预加载一次（启动日志即可发现数据缺失问题；后续由 _ensure_cn_holidays_fresh 热重载）
_load_cn_holidays()

def _extract_attachment_text(name: str, b64: str):
    """Parse attachment bytes into text for LLM (pdf/docx/xlsx/txt). Returns None when unsupported/unparsable."""
    import base64
    import io
    ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
    try:
        raw = base64.b64decode(b64 or '')
    except Exception:
        return None
    if not raw:
        return None
    try:
        if ext == 'pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            texts = []
            for page in reader.pages[:20]:
                t = (page.extract_text() or '').strip()
                if t:
                    texts.append(t)
            return ("\n\n".join(texts)[:20000]) if texts else "(PDF has no text layer)"
        if ext == 'docx':
            from docx import Document
            doc = Document(io.BytesIO(raw))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            return ("\n".join(paras)[:20000]) if paras else "(docx has no text)"
        if ext == 'xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:5]:
                lines.append(f"--- sheet: {ws.title} ---")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append("\t".join(cells)[:500])
                if len(lines) > 300:
                    lines.append("...(truncated)")
                    break
            return ("\n".join(lines)[:20000]) if lines else "(xlsx is empty)"
        if ext in ('txt', 'md', 'log', 'csv', 'json', 'ini', 'conf', 'cfg', 'yaml', 'yml', 'toml'):
            try:
                return raw.decode('utf-8')[:20000]
            except UnicodeDecodeError:
                try:
                    return raw.decode('gbk', errors='replace')[:20000]
                except Exception:
                    return None
        return None
    except Exception:
        return None

class AgentChatServiceV2:
    """
    Agent 聊天服务 V2
    
    使用状态机模式重构的核心服务
    
    完整功能：
    - 状态机循环（while true + Continue/Terminal）
    - 工具智能并发
    - 上下文压缩（自动/响应式）
    - 记忆系统
    - 标题自动生成
    - Transcript 管理
    - Todo 回调
    - Ask Question 阻塞
    - 权限检查
    """
    
    def __init__(self):
        # 取消信号（asyncio.Event，用于 MCP 客户端的 await cancel_event.wait()）
        self._cancel_signals: Dict[str, asyncio.Event] = {}
        
        # 当前项目路径
        self._current_project_path: str = os.getcwd()
        
        # 是否需要生成标题
        self._need_generate_title: bool = False
        
        # 当前对话消息计数
        self._conversation_message_count: int = 0
        
        # Transcript 管理器
        self._transcript_manager = get_transcript_manager()
        
        # 经验账本
        self._ledger = get_ledger()
        
        # 工具执行器
        self._tool_executor = AgentToolExecutorImpl(self._current_project_path)
        
        # MCP 管理器
        self._mcp_manager = None
        self._mcp_initialized = False
        
        # 计费系统
        self._ensure_usage()
    
    def _ensure_usage(self) -> UsageTracker:
        if not hasattr(self, "_usage"):
            self._usage = UsageTracker(
                _DATA_HOME,
                os.path.dirname(os.path.dirname(__file__)),
                current_pricing_period,
            )
        return self._usage

    @property
    def _model_pricing(self):
        return self._ensure_usage().pricing

    @_model_pricing.setter
    def _model_pricing(self, value):
        self._ensure_usage().pricing = value

    @property
    def _model_pricing_mtime(self):
        return self._ensure_usage().pricing_mtime

    @_model_pricing_mtime.setter
    def _model_pricing_mtime(self, value):
        self._ensure_usage().pricing_mtime = value

    def _load_model_pricing(self) -> dict:
        return self._ensure_usage().reload()

    def _ensure_model_pricing_fresh(self) -> None:
        self._ensure_usage().ensure_fresh()

    def _calculate_cost(self, model: str, usage) -> Optional[dict]:
        return self._ensure_usage().calculate_cost(model, usage)

    @staticmethod
    def _cache_hit_tokens(usage) -> int:
        """缓存命中 token：兼容两种字段约定。

        DeepSeek 用私有字段 prompt_cache_hit_tokens；OpenAI 兼容网关
        （OpenCode Go 等）用标准字段 prompt_tokens_details.cached_tokens。
        """
        hit = getattr(usage, "prompt_cache_hit_tokens", 0) or 0
        if hit:
            return int(hit)
        details = getattr(usage, "prompt_tokens_details", None)
        cached = getattr(details, "cached_tokens", 0) if details is not None else 0
        return int(cached or 0)

    def _unpriced_cost_info(self, usage) -> dict:
        """无价格条目的模型（订阅制 provider）：上报真实用量，费用恒为 0。

        返回结构与 UsageTracker.calculate_cost 一致，使累计、持久化和前端
        展示走同一条路径；token 数为 API 原始值，上下文占用由此恢复。
        """
        prompt = getattr(usage, "prompt_tokens", 0) or 0
        completion = getattr(usage, "completion_tokens", 0) or 0
        cache_hit = min(self._cache_hit_tokens(usage), prompt)
        return {
            "input_tokens": prompt - cache_hit,
            "cache_hit_tokens": cache_hit,
            "output_tokens": completion,
            "input_cost": 0.0,
            "cache_hit_cost": 0.0,
            "output_cost": 0.0,
            "total_cost": 0.0,
            "period": None,
            "pricing_used": None,
        }

    def _accumulate_cost(self, conversation_uuid: str, cost_info: dict, model_name: str = None) -> dict:
        return self._ensure_usage().accumulate(conversation_uuid, cost_info, model_name)

    def _load_session_cost_from_db(self, conversation_uuid: str) -> Optional[dict]:
        return self._ensure_usage().load_session_cost(conversation_uuid)
    
    def _resolve_path(self, path: str) -> str:
        """解析路径，支持相对路径"""
        if not path:
            return path
        if os.path.isabs(path):
            return path
        return os.path.join(self._current_project_path, path)
    
    def set_project_path(self, path: str):
        """更新工作区目录，同步到所有相关组件"""
        self._current_project_path = path
        self._tool_executor.set_project_path(path)
        try:
            from .truncation_manager import set_truncation_dir
            set_truncation_dir(path)
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Failed to set truncation dir: {e}")
        self._load_path_whitelist()
        if hasattr(self, '_mcp_manager') and self._mcp_manager:
            from .mcp.mcp_config import get_mcp_config_manager
            try:
                config_manager = get_mcp_config_manager()
                config_manager.set_project_path(path)
            except Exception as e:
                logger.warning(f"[AgentChatServiceV2] Failed to set MCP project path: {e}")

    def _load_path_whitelist(self) -> list:
        """从 path_whitelist.json 加载用户配置的白名单路径并同步到执行器"""
        user_paths = []
        data = _read_json_config("path_whitelist.json")
        if data:
            user_paths = data.get("paths", [])
        # 合并 agent_model_config.json 中的 allowed_paths
        merged = list(user_paths)
        try:
            agent_config = self._load_agent_model_config(
                getattr(self, '_current_agent', 'pebble')
            ) or {}
            allowed = agent_config.get("allowed_paths", [])
            for p in allowed:
                if p not in merged:
                    merged.append(p)
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Failed to load agent allowed_paths: {e}")
        if merged:
            self._tool_executor.set_path_whitelist(merged)
        return merged

    async def _ensure_mcp_initialized(self, project_path: Optional[str] = None):
        """
        确保 MCP 管理器已初始化
        
        在第一次聊天时初始化 MCP 连接。
        后续调用会跳过。
        """
        if self._mcp_initialized:
            return
        
        try:
            from .mcp.mcp_manager import get_mcp_manager
            from .mcp.mcp_config import get_mcp_config_manager
            
            # 设置项目路径
            if project_path:
                config_manager = get_mcp_config_manager()
                config_manager.set_project_path(project_path)
            
            # 获取或创建 MCP 管理器
            self._mcp_manager = get_mcp_manager()
            
            if not self._mcp_manager._initialized:
                logger.info("[AgentChatServiceV2] Initializing MCP manager...")
                await self._mcp_manager.initialize()
                logger.info(f"[AgentChatServiceV2] MCP initialized: "
                           f"{len(self._mcp_manager._clients)} servers, "
                           f"{len(self._mcp_manager._tools)} tools")
            
            self._mcp_initialized = True
            
        except ImportError as e:
            logger.warning(f"[AgentChatServiceV2] MCP not available: {e}")
            self._mcp_initialized = False
        except Exception as e:
            logger.error(f"[AgentChatServiceV2] Failed to initialize MCP: {e}")
            self._mcp_initialized = False
    
    async def generate_conversation_title(
        self,
        messages: List[Dict[str, str]],
        model: str | None = None,
        thinking_enabled: bool = True,
        reasoning_effort: str = "high",
        conversation_uuid: str = "",
    ) -> str:
        """根据对话内容生成标题"""
        if not messages or len(messages) < 2:
            return "新对话"
        
        llm_config = self._get_llm_config("code_agent", model)
        use_model = model or llm_config['model']
        
        # 取前几轮对话
        context_messages = messages[:min(6, len(messages))]
        
        conversation_summary = ""
        for msg in context_messages:
            role = "用户" if msg.get("role") == "user" else "AI"
            content = (msg.get("content") or "")[:200]
            conversation_summary += f"{role}: {content}\n"
        
        title_prompt = f"""请根据以下对话内容，生成一个简洁的标题（不超过20个字符）。

对话内容:
{conversation_summary}

要求:
1. 标题要简洁，不超过20个字符
2. 概括对话的主要主题或任务
3. 不要包含标点符号
4. 直接输出标题，不要有其他内容

标题:"""
        
        try:
            client = AsyncOpenAI(
                api_key=llm_config['api_key'],
                base_url=llm_config['base_url'],
                timeout=30.0,
                default_headers=build_default_headers(llm_config, conversation_uuid),
            )
            
            title_params = {
                "model": use_model,
                "messages": [{"role": "user", "content": title_prompt}],
                "temperature": 0.3,
                "max_tokens": 200,  # 标题 ≤20 字
                "stream": True,
            }
            
            # 思考模式参数（与主聊天一致：开关 + 强度，非 DeepSeek 的 max 档自动降级）
            title_params.update(build_thinking_params(use_model, thinking_enabled, reasoning_effort, keep_max_effort(llm_config)))
            apply_reasoning_mode(title_params, llm_config)
            
            raw_content = ""
            async for chunk in await client.chat.completions.create(**title_params):
                if chunk.choices and chunk.choices[0].delta.content:
                    raw_content += chunk.choices[0].delta.content
            if not raw_content:
                return "新对话"
            title = raw_content.strip()
            title = title.replace("标题:", "").strip()
            title = title.replace("：", "").strip()
            title = title.replace(":", "").strip()
            
            if len(title) > 30:
                title = title[:30]
            
            return title or "新对话"
            
        except Exception as e:
            logger.error(f"[AgentChatServiceV2] Failed to generate title: {e}")
            return "新对话"

    def _load_agent_model_config(self, agent_name: str) -> Optional[Dict[str, Any]]:
        """加载独立 Agent 聊天的模型配置"""
        config = _read_json_config("agent_model_config.json")
        if not config:
            return None
        return config.get("agents", {}).get(agent_name)

    def _find_provider_for_model(self, model_name: str) -> Optional[Dict[str, Any]]:
        """根据模型名查找 provider 配置"""
        try:
            config = _read_json_config("llm_providers.json")
            if not config:
                return None

            for provider in config.get("providers", []):
                models = provider.get("models", [])
                if model_name in models:
                    base_url = provider.get("baseUrl", "")
                    if base_url.endswith("/chat/completions"):
                        base_url = base_url.rsplit("/chat/completions", 1)[0]
                    return {
                        "provider": provider.get("provider"),
                        "api_key": provider.get("apiKey"),
                        "base_url": base_url,
                        "model": model_name,
                        "temperature": provider.get("temperature"),
                        "top_p": provider.get("topP"),
                    }
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Failed to find provider for model {model_name}: {e}")
        return None

    def _get_provider_model(self, source: str) -> Optional[str]:
        """按 provider 索引获取默认模型"""
        try:
            if not source.startswith("provider_"):
                return None
            provider_idx = int(source.split("_", 1)[1])
            config = _read_json_config("llm_providers.json")
            if not config:
                return None

            providers = config.get("providers", [])
            if 0 <= provider_idx < len(providers):
                models = providers[provider_idx].get("models", [])
                if models:
                    return models[0]
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Failed to resolve provider model from {source}: {e}")
        return None

    def _resolve_agent_model(self, agent_name: str, requested_model: Optional[str], agent_def) -> Optional[str]:
        """解析最终使用的模型，优先尊重显式请求，其次读取 agent_model_config.json"""
        if requested_model:
            return requested_model

        agent_config = self._load_agent_model_config(agent_name)
        if agent_config:
            source = agent_config.get("model_source", "inherit")
            if source == "custom":
                custom_model = agent_config.get("custom_model")
                if custom_model:
                    return custom_model
            elif source and source != "inherit":
                if source.startswith("provider_"):
                    provider_model = self._get_provider_model(source)
                    if provider_model:
                        return provider_model
                else:
                    return source

        if getattr(agent_def, "model", "inherit") != "inherit":
            return agent_def.model
        return None


    
    def _get_llm_config(self, agent_name: str, model: str = None) -> Dict[str, Any]:
        """获取 LLM 配置
        
        Args:
            agent_name: Agent 名称
            model: 指定的模型名称（可选）
        """
        requested_model = model

        try:
            config = _read_json_config("llm_providers.json")
            if config:
                providers = config.get('providers', [])
                if providers:
                    agent_def = get_cached_agent(agent_name)
                    resolved_model = self._resolve_agent_model(agent_name, requested_model, agent_def) if agent_def else requested_model

                    # 如果指定了模型，找到包含该模型的 provider
                    selected_provider = None
                    if resolved_model:
                        for provider in providers:
                            provider_models = provider.get('models', [])
                            if resolved_model in provider_models:
                                selected_provider = provider
                                break
                    
                    # 如果没找到或没指定模型，使用第一个 provider
                    if not selected_provider:
                        selected_provider = providers[0]
                    
                    base_url = selected_provider.get('baseUrl', '')
                    if base_url.endswith('/chat/completions'):
                        base_url = base_url.rsplit('/chat/completions', 1)[0]

                    fallback_models = selected_provider.get('models', ['gpt-4'])
                    final_model = resolved_model or (fallback_models[0] if fallback_models else 'gpt-4')
                    return {
                        'api_key': selected_provider.get('apiKey', os.getenv("OPENAI_API_KEY", "local-dev-key")),
                        'base_url': base_url or os.getenv("OPENAI_BASE_URL", "http://localhost:8080/v1"),
                        'model': final_model,
                        'provider': selected_provider.get('provider'),
                        'temperature': selected_provider.get('temperature'),
                        'top_p': selected_provider.get('topP'),
                        # Request shaping (headers / reasoning mode / max tier) must
                        # travel with the provider the model belongs to.
                        **shaping_from_provider(selected_provider),
                    }
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Failed to load LLM config: {e}")
        
        return {
            'api_key': os.getenv("OPENAI_API_KEY", "local-dev-key"),
            'base_url': os.getenv("OPENAI_BASE_URL", "http://localhost:8080/v1"),
            'model': requested_model or os.getenv("OPENAI_MODEL", "gpt-4"),
            'provider': None,
            'temperature': None,
            'top_p': None,
            'extra_headers': {},
            'session_header': '',
            'reasoning_mode': 'thinking',
            'max_effort': 'downgrade',
        }
    
    def get_available_agents(self) -> List[Dict[str, Any]]:
        """获取可用的 Agent 列表"""
        agents = list_available_agents_from_yaml()
        
        result = []
        for agent_info in agents:
            agent_def = get_cached_agent(agent_info['name'])
            if agent_def:
                result.append({
                    "name": agent_def.name,
                    "description": agent_def.description,
                    "model": agent_def.model,
                    "skills": agent_def.skills
                })
        
        return result
    
    def get_agent_info(self, agent_name: str) -> Optional[Dict[str, Any]]:
        """获取 Agent 信息"""
        agent_def = get_cached_agent(agent_name)
        if agent_def:
            return {
                "name": agent_def.name,
                "description": agent_def.description,
                "model": agent_def.model,
                "temperature": agent_def.temperature,
                "skills": agent_def.skills,
                "context_strategy": agent_def.context_strategy.type
            }
        return None
    
    def cancel_chat(self, conversation_uuid: str) -> bool:
        """取消聊天 - 设置取消信号"""
        if conversation_uuid in self._cancel_signals:
            self._cancel_signals[conversation_uuid].set()
            logger.info(f"[AgentChatServiceV2] Cancelled chat: {conversation_uuid}")
            return True
        return False
    
    def _sanitize_messages(self, messages: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        兜底修复：确保发送给 API 的消息完全合法
        
        1. 移除孤儿 tool 消息（没有对应 assistant tool_calls 的）
        2. assistant 消息中的 tool_calls 必须有 id
        """
        # 第一遍：收集所有 assistant 中 tool_calls 声明的 id
        valid_tool_call_ids = set()
        for msg in messages:
            if msg.get("role") == "assistant" and msg.get("tool_calls"):
                for tc in msg["tool_calls"]:
                    tc_id = tc.get("id")
                    if tc_id:
                        valid_tool_call_ids.add(tc_id)
        
        # 第二遍：过滤
        sanitized = []
        removed_count = 0
        for i, msg in enumerate(messages):
            role = msg.get("role", "")
            
            # 孤儿 tool 消息：tool_call_id 不在任何一个 assistant 的 tool_calls 中
            if role == "tool":
                tcid = msg.get("tool_call_id")
                if tcid and tcid not in valid_tool_call_ids:
                    logger.debug(f"[SANITIZE] Removing orphan tool message[{i}]: tool_call_id={tcid} has no matching assistant tool_calls")
                    removed_count += 1
                    continue
                if not tcid:
                    logger.debug(f"[SANITIZE] Removing tool message[{i}]: missing tool_call_id")
                    removed_count += 1
                    continue
            
            # assistant 消息的 tool_calls 必须每个都有 id
            # 构造副本，不原地修改入参
            if role == "assistant" and msg.get("tool_calls"):
                if any(not tc.get("id") for tc in msg["tool_calls"]):
                    new_tool_calls = []
                    for j, tc in enumerate(msg["tool_calls"]):
                        if not tc.get("id"):
                            tc = {**tc, "id": f"fallback_tc_{i}_{j}"}
                            logger.debug(f"[SANITIZE] messages[{i}] tool_calls[{j}] missing id, generated: {tc['id']}")
                        new_tool_calls.append(tc)
                    msg = {**msg, "tool_calls": new_tool_calls}
            
            sanitized.append(msg)
        
        if removed_count > 0:
            logger.debug(f"[SANITIZE] Removed {removed_count} orphan/invalid tool messages")
        
        return sanitized
    
    def _build_and_sanitize_tool_calls(self, tool_calls_chunks: Dict, tools: List[Dict[str, Any]]) -> Optional[List[Dict[str, Any]]]:
        """从流式 chunks 构建工具调用列表并清洗

        1. 空 id 修复：流式返回中 tool_call.id 可能因网络波动/chunk 乱序/取消时机
           而保持空字符串。统一生成稳定 UUID，确保下游消息构建、execution_order
           记录、数据库保存、取消后重建全链路 id 永远非空。
        2. 清洗：空名/未知名工具调用不是可恢复的真实工具调用；如果把它们回灌给
           模型，会形成"空工具名 -> 错误反馈 -> 再次尝试空工具名"的循环，故丢弃。
        3. 工具名大小写修复。
        """
        if not tool_calls_chunks:
            return None
        tool_calls = [tool_calls_chunks[idx] for idx in sorted(tool_calls_chunks.keys())]
        
        # ========== 空 id 修复 ==========
        for tc in tool_calls:
            if not tc.get("id"):
                tc["id"] = f"call_{uuid.uuid4().hex[:16]}"
                logger.warning(
                    f"[AgentChatServiceV2] Fixed empty tool_call.id -> {tc['id']} "
                    f"(function={tc.get('function', {}).get('name', 'unknown')})"
                )

        # 有效工具名表（含小写索引）
        valid_tool_names = {}
        for tool_def in tools:
            tool_name = tool_def.get("function", {}).get("name", "")
            if tool_name and tool_name != "_unknown_tool":
                valid_tool_names[tool_name] = tool_name
                valid_tool_names[tool_name.lower()] = tool_name

        sanitized_tool_calls = []
        dropped_tool_calls = []
        for tc in tool_calls:
            function = tc.get("function", {})
            tool_name = function.get("name", "")

            if not tool_name:
                dropped_tool_calls.append((tc.get("id", ""), "empty_name"))
                continue

            corrected = valid_tool_names.get(tool_name) or valid_tool_names.get(tool_name.lower())
            if not corrected:
                dropped_tool_calls.append((tc.get("id", ""), f"unknown_name:{tool_name}"))
                continue

            if corrected != tool_name:
                logger.info(f"[AgentChatServiceV2] Repaired tool name case: '{tool_name}' -> '{corrected}'")
                tc["function"]["name"] = corrected

            sanitized_tool_calls.append(tc)

        if dropped_tool_calls:
            logger.warning(
                "[AgentChatServiceV2] Dropped invalid tool calls before execution: %s",
                dropped_tool_calls
            )
            tool_calls = sanitized_tool_calls
        
        return tool_calls
    
    async def _extract_memories_async(
        self,
        conversation_uuid: str,
        messages: List[Dict[str, Any]],
        llm_config: Dict[str, Any] = None,
        tool_events: List[Dict[str, Any]] = None,
        message_id: str = None,
    ) -> None:
        """把本轮对话中的经验信号写入账本（规则检测，无额外模型调用）。"""
        try:
            user_message = ""
            assistant_message = ""
            for msg in reversed(messages):
                role = msg.get("role")
                if role == "user" and not user_message:
                    user_message = msg.get("content") or ""
                elif role == "assistant" and not assistant_message:
                    assistant_message = msg.get("content") or ""
                if user_message and assistant_message:
                    break

            created = self._ledger.observe_turn(
                conversation_uuid=conversation_uuid,
                user_message=user_message,
                assistant_message=assistant_message,
                tool_events=tool_events or [],
                message_id=message_id,
            )
            if created:
                logger.info(f"[AgentChatServiceV2] Recorded {len(created)} experience entries")
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Experience recording error: {e}")
    
    async def chat_stream(
        self,
        agent_name: str,
        message: str,
        conversation_uuid: str,
        user_id: str = "default_user",
        model: str | None = None,
        attachments: List[Dict[str, Any]] | None = None,
        enable_thinking: bool = False,
        project_path: str | None = None,
        temperature: float | None = None,
        top_p: float | None = None,
        thinking_enabled: bool | None = None,
        reasoning_effort: str | None = None,
        user_msg_uuid: str | None = None
    ) -> AsyncGenerator[Dict[str, Any], None]:
        """
        流式聊天（状态机版本）
        
        完整功能：
        - 状态机循环
        - 工具智能并发
        - 上下文压缩
        - 记忆系统
        - 标题自动生成
        - Todo 回调
        - Ask Question 阻塞
        - 权限检查
        """
        logger.info(f"[AgentChatServiceV2] Starting chat with {agent_name}, temperature={temperature}, top_p={top_p}, thinking_enabled={thinking_enabled}, reasoning_effort={reasoning_effort}")
        
        # 存储项目路径
        self._current_project_path = project_path or os.getcwd()  # project_path 由前端传入，此处仅兜底
        
        # 设置经验账本的项目作用域
        if project_path:
            self._ledger.set_project(project_path)
        
        # ========== MCP 初始化 ==========
        await self._ensure_mcp_initialized(project_path)
        
        # 1-11. 准备本轮状态
        state = TurnState()
        async for ev in prepare_turn(self, {
            "agent_name": agent_name,
            "message": message,
            "conversation_uuid": conversation_uuid,
            "user_id": user_id,
            "model": model,
            "attachments": attachments,
            "project_path": project_path,
            "temperature": temperature,
            "top_p": top_p,
            "thinking_enabled": thinking_enabled,
            "reasoning_effort": reasoning_effort,
            "user_msg_uuid": user_msg_uuid,
        }, state):
            yield ev
        if state.aborted:
            return

        agent_def = state.agent_def
        conversation_id = state.conversation_id
        messages = state.messages
        use_model = state.use_model
        final_temperature = state.final_temperature
        final_top_p = state.final_top_p
        final_thinking_enabled = state.final_thinking_enabled
        final_reasoning_effort = state.final_reasoning_effort
        client = state.client
        cancel_event = state.cancel_event
        user_message_id = state.user_message_id
        pending_todo_events = state.pending_todo_events
        full_message = state.full_message
        is_deepseek = 'deepseek' in (use_model or '').lower()
        # 12. 发送开始事件
        yield {
            "type": "chat_started",
            "agent_name": agent_name,
            "conversation_uuid": conversation_uuid,
            "user_message_id": user_message_id,
        }
        
        # ==================== 状态机主循环 ====================
        
        try:
            final_response = ""
            tool_executions = []
            all_response_text = ""
            all_reasoning_text = ""  # 累积所有迭代的推理内容
            execution_order = []
            order_index = 0  # 全局顺序索引，确保严格按输出顺序
            was_interrupted = False  # 是否被用户中断
            suspicious_retry_done = False  # C: 每回合最多自动重试一次（可疑结束检测）
            in_suspicious_retry = False  # 当前 iteration 是否为可疑重试轮（API 兜底保护窗口）
            
            logger.info("[AgentChatServiceV2] Starting state machine loop (unlimited iterations)")
            
            # ========== while(true) 主循环 ==========
            iteration = 0
            while True:
                iteration += 1
                current_iteration = iteration
                
                logger.info(f"[AgentChatServiceV2] ========== Iteration {current_iteration} ==========")
                
                # 检查取消
                if cancel_event.is_set():
                    was_interrupted = True
                    yield {"type": "cancelled", "message": "Generation cancelled"}
                    break
                
                # 重置当前迭代状态
                response_text = ""
                reasoning_text = ""  # 思考模式下的推理内容
                finish_reason = None  # B: 本轮流的结束方式（stop/length/tool_calls/None=异常缺失）
                tool_calls_chunks = {}

                # 统一上下文口径：历史轮的 tool 输出与 DB 重载路径一致地
                # 截断到 MAX_TOOL_OUTPUT_CONTEXT_CHARS；当前轮结果在 append 时保持全文
                for _m in messages:
                    if _m.get("role") == "tool" and isinstance(_m.get("content"), str):
                        _m["content"] = history._truncate_context_text(
                            _m["content"], history.MAX_TOOL_OUTPUT_CONTEXT_CHARS)
                
                # ========== 调用 LLM ==========
                tools = self._get_tools_definition(agent_def)
                
                stream_params = {
                    "model": use_model,
                    "messages": messages,
                    "stream": True,
                    "temperature": final_temperature,
                    "top_p": final_top_p,
                }
                
                # 思考模式参数（所有厂商统一：开关 + 强度；非 DeepSeek 的 max 档自动降级）
                stream_params.update(build_thinking_params(use_model, final_thinking_enabled, final_reasoning_effort, keep_max_effort(state.llm_config)))
                apply_reasoning_mode(stream_params, state.llm_config)
                
                if tools:
                    stream_params["tools"] = tools
                    # DeepSeek 思考模式不支持 tool_choice 参数
                    if not (is_deepseek and final_thinking_enabled):
                        stream_params["tool_choice"] = "auto"
                
                logger.debug(f"[AgentChatServiceV2] Final stream_params keys: {list(stream_params.keys())}")
                
                # 调用 LLM（含重试与上下文超长处理）
                stream, open_status, graceful_abort = await open_stream(
                    client, stream_params, cancel_event, in_suspicious_retry
                )
                if open_status == "cancelled":
                    was_interrupted = True
                    yield {"type": "cancelled", "message": "Generation cancelled during API retry"}
                    break
                in_suspicious_retry = False
                if stream is None:
                    if graceful_abort:
                        break  # 跳出 iteration 主循环 → 走正常收尾（接受为完成）
                    raise RuntimeError("Failed to get API response after retries")
                
                # 流式响应处理
                consumer = StreamConsumer(execution_order, order_index, cancel_event)
                async for ev in consumer.consume(stream):
                    yield ev
                was_interrupted = consumer.interrupted
                response_text = consumer.content
                reasoning_text = consumer.reasoning
                finish_reason = consumer.finish_reason
                tool_calls_chunks = consumer.tool_calls_chunks
                order_index = consumer.order_index

                # 先累加本轮的 response_text 和 reasoning_text（中断前也要保存）
                if reasoning_text:
                    all_reasoning_text += reasoning_text
                if response_text:
                    all_response_text += response_text
                    final_response = response_text

                # B: 记录本轮流的结束方式（stop=正常 / length=截断 / None=异常缺失，可事后追查）
                logger.info(
                    f"[AgentChatServiceV2] Iteration {iteration} stream end: "
                    f"finish_reason={finish_reason}, content_len={len(response_text)}, "
                    f"reasoning_len={len(reasoning_text)}, tool_calls={len(tool_calls_chunks)}"
                )

                # 中断检查：流式结束后立即检查，避免继续执行工具
                if was_interrupted:
                    logger.info("[AgentChatServiceV2] Interrupted after streaming, breaking main loop")
                    break
                
                # ========== 捕获 usage 并计费（仅使用 API 原始 token 值） ==========
                if consumer.usage:
                    try:
                        cost_info = self._calculate_cost(use_model, consumer.usage)
                        if cost_info is None:
                            # 无价格条目（订阅制 provider）：仍上报用量，费用记 0。
                            # token 统计、上下文占用与会话历史不因缺少定价而中断。
                            cost_info = self._unpriced_cost_info(consumer.usage)
                        session_usage = self._accumulate_cost(conversation_uuid, cost_info, model_name=use_model)
                        # 环显示用完整 prompt_tokens（已含缓存），面板用 cost_info.input_tokens（非缓存输入）
                        latest_input = getattr(consumer.usage, 'prompt_tokens', 0)

                        yield {
                            "type": "usage",
                            "model": use_model,
                            "session_cost_cny": round(session_usage["total_cost"], 4),
                            "latest_input_tokens": latest_input,
                            "latest_output_tokens": cost_info.get("output_tokens", 0),
                            "model_context_limit": get_model_context_limit(use_model),
                            "period": cost_info.get("period"),  # 本次计费时段 peak/offpeak（分时模型）
                            "this_call": {**cost_info, "model": use_model}
                        }
                    except Exception as e:
                        logger.warning(f"[Cost] Failed to calculate cost: {e}", exc_info=True)
                # ========================================
                
                # 构建工具调用列表（清洗逻辑见 _build_and_sanitize_tool_calls）
                tool_calls = self._build_and_sanitize_tool_calls(tool_calls_chunks, tools)
                # ========================================
                
                # ========== 判断是否终止 ==========
                if not tool_calls:
                    # 可疑结束检测 + 自动重试一次：
                    # 模型在"思考→行动"衔接点静默停止时，不能直接判定为"任务完成"
                    suspicious = _is_suspicious_stream_end(finish_reason, response_text)
                    if suspicious and not suspicious_retry_done:
                        suspicious_retry_done = True
                        in_suspicious_retry = True  # 标记下一 iteration 为兜底保护窗口
                        logger.warning(
                            f"[AgentChatServiceV2] Suspicious stream end (finish_reason={finish_reason}, "
                            f"content_len={len(response_text)}, reasoning_len={len(reasoning_text)}) - auto-retrying once"
                        )
                        # 回传本轮输出保持上下文连贯（DeepSeek 思考模式要求回传 reasoning_content）
                        # content 为空时不得回传——空 content + 无 tool_calls 的 assistant
                        # 消息会被 API 拒绝（400 Invalid assistant message），跳过以保持请求合法
                        if response_text:
                            messages.append({
                                "role": "assistant",
                                "content": response_text,
                                "reasoning_content": reasoning_text or None,
                            })
                        messages.append({
                            "role": "user",
                            "content": "[系统提示] 你上一条输出意外中断了（没有工具调用也没有文字收尾）。请从中断处继续完成任务。",
                        })
                        continue
                    if suspicious:
                        logger.warning(
                            f"[AgentChatServiceV2] Suspicious stream end persisted after retry "
                            f"(finish_reason={finish_reason}) - accepting as complete"
                        )
                    logger.info("[AgentChatServiceV2] No tool calls, task completed")
                    break
                
                # ========== 执行工具（智能并发） ==========
                logger.info(f"[AgentChatServiceV2] Executing {len(tool_calls)} tool calls")

                messages.append({
                    "role": "assistant",
                    "content": response_text or None,
                    "tool_calls": tool_calls,
                    "reasoning_content": reasoning_text if reasoning_text else None,
                })

                cycle = ToolCycle(
                    executor=self._tool_executor,
                    cancel_event=cancel_event,
                    execution_order=execution_order,
                    order_index=order_index,
                    messages=messages,
                    conversation_uuid=conversation_uuid,
                    agent_name=agent_name,
                    project_path=self._current_project_path,
                    user_message_id=user_message_id,
                    tool_executions=tool_executions,
                    pending_todo_events=pending_todo_events,
                )
                async for ev in cycle.prepare(tool_calls):
                    yield ev
                async for ev in cycle.run(tool_calls):
                    yield ev
                order_index = cycle.order_index
                was_interrupted = cycle.interrupted

                if was_interrupted:
                    break
            # ========== 循环结束，保存结果 ==========
            
            # 序列化 execution_order 为 JSON
            execution_order_json = json.dumps(execution_order, ensure_ascii=False) if execution_order else None
            
            # 保存 AI 回复 — 无论是否中断都必须执行
            assistant_message_id = 0
            assistant_msg_uuid = str(uuid.uuid4())
            if execution_order:
                # 只提取 text 类型内容，不包含 tool 类型
                full_response_parts = []
                for item in execution_order:
                    if item["type"] == "text":
                        full_response_parts.append(item["content"])
                    # tool 类型不添加到 content 中，避免 LLM 学习并伪造

                full_response = "\n\n".join(full_response_parts) if full_response_parts else ""
                # 修复：execution_order 无 text 项但有实际响应文本时，用 all_response_text 兜底
                if not full_response and (all_response_text or final_response):
                    full_response = all_response_text or final_response

                assistant_message_id = history.add_agent_chat_message(
                    conversation_id=conversation_id,
                    role="assistant",
                    content=full_response,
                    reasoning_content=all_reasoning_text if all_reasoning_text else None,
                    token_count=count_tokens(full_response),
                    agent_name=agent_name,
                    execution_order=execution_order_json,
                    msg_uuid=assistant_msg_uuid
                )
            elif all_response_text or final_response:
                saved_response = all_response_text if all_response_text else final_response
                assistant_message_id = history.add_agent_chat_message(
                    conversation_id=conversation_id,
                    role="assistant",
                    content=saved_response,
                    reasoning_content=all_reasoning_text if all_reasoning_text else None,
                    token_count=count_tokens(saved_response),
                    agent_name=agent_name,
                    execution_order=execution_order_json,
                    msg_uuid=assistant_msg_uuid
                )
            elif was_interrupted:
                # 中断时且没有任何文本/工具调用的极端情况：存一条空消息
                assistant_message_id = history.add_agent_chat_message(
                    conversation_id=conversation_id,
                    role="assistant",
                    content="",
                    reasoning_content=all_reasoning_text if all_reasoning_text else None,
                    token_count=0,
                    agent_name=agent_name,
                    execution_order=execution_order_json,
                    msg_uuid=assistant_msg_uuid
                )
            
            # 非中断时才生成标题和提取记忆
            if not was_interrupted:
                # 检查是否需要生成标题
                current_messages = history.get_agent_chat_messages_for_context(conversation_id, limit=10)
                message_count = len(current_messages)
                
                conv_info = history.get_agent_conversation_by_uuid(conversation_uuid)
                current_title = conv_info.get("title", "") if conv_info else ""
                
                need_generate_title = False
                if message_count >= 3:
                    if current_title == "新对话" or current_title.startswith("与 "):
                        need_generate_title = True
                
                if need_generate_title:
                    try:
                        new_title = await self.generate_conversation_title(current_messages, use_model, final_thinking_enabled, final_reasoning_effort, conversation_uuid)
                        history.update_agent_conversation_title(conversation_uuid, new_title)
                        yield {
                            "type": "title_updated",
                            "conversation_uuid": conversation_uuid,
                            "title": new_title
                        }
                    except Exception as e:
                        logger.error(f"[AgentChatServiceV2] Failed to generate title: {e}")
            
            # 发送完成事件（无论是否中断都发，让前端能收到完成信号）
            yield {
                "type": "chat_completed",
                "content": all_response_text if all_response_text else final_response,
                "reasoning": all_reasoning_text if all_reasoning_text else None,
                "agent_name": agent_name,
                "execution_order": execution_order,
                "assistant_message_id": assistant_message_id,
                "assistant_msg_uuid": assistant_msg_uuid,
                "interrupted": was_interrupted
            }
            
            # 非中断时记录经验信号
            if not was_interrupted:
                try:
                    tool_events = []
                    for item in execution_order or []:
                        if not isinstance(item, dict) or item.get("type") != "tool":
                            continue
                        info = item.get("tool_info") or item.get("toolInfo") or {}
                        args = info.get("input")
                        if isinstance(args, str):
                            try:
                                args = json.loads(args)
                            except Exception:
                                args = {}
                        tool_events.append({
                            "name": info.get("name", ""),
                            "status": info.get("status", ""),
                            "args": args if isinstance(args, dict) else {},
                            "output": info.get("output", ""),
                        })
                    await self._extract_memories_async(
                        conversation_uuid=conversation_uuid,
                        messages=[
                            {"role": "user", "content": full_message},
                            {"role": "assistant", "content": all_response_text or final_response or ""},
                        ],
                        tool_events=tool_events,
                        message_id=str(state.user_message_id) if state.user_message_id else None,
                    )
                except Exception as e:
                    logger.warning(f"[AgentChatServiceV2] Experience recording failed: {e}")
        
        except asyncio.CancelledError:
            # future.cancel() 触发的 CancelledError：保存当前已累积的数据到 agent_chat.db
            logger.warning("[AgentChatServiceV2] CancelledError caught, saving partial data")
            was_interrupted = True
            execution_order_json = json.dumps(execution_order, ensure_ascii=False) if execution_order else None
            cancelled_msg_uuid = str(uuid.uuid4())
            # 关键：包含当前迭代已生成但未累加的 response_text / reasoning_text
            partial_content = (all_response_text or "") + (response_text or "")
            partial_reasoning = (all_reasoning_text or "") + (reasoning_text or "")
            logger.debug(f"[DEBUG-INTERRUPT] Service CancelledError: all_response={len(all_response_text or '')} response={len(response_text or '')}  partial_content={len(partial_content or '')} partial_reasoning={len(partial_reasoning or '')} all_reasoning={len(all_reasoning_text or '')} reasoning={len(reasoning_text or '')}")
            if execution_order or partial_content or partial_reasoning or final_response:
                try:
                    history.add_agent_chat_message(
                        conversation_id=conversation_id,
                        role="assistant",
                        content=partial_content,
                        reasoning_content=partial_reasoning or None,
                        token_count=count_tokens(partial_content),
                        agent_name=agent_name,
                        execution_order=execution_order_json,
                        msg_uuid=cancelled_msg_uuid
                    )
                    logger.info(f"[AgentChatServiceV2] Partial data saved to agent_chat.db after CancelledError, uuid={cancelled_msg_uuid}")
                except Exception as e:
                    logger.error(f"[AgentChatServiceV2] Failed to save partial data after CancelledError: {e}")
            # 发射 chat_completed 事件，携带 interrupted=true，让前端能收到中断标记
            logger.debug(f"[DEBUG-INTERRUPT] Service yielding chat_completed: content_len={len(partial_content or '')} reasoning_len={len(partial_reasoning or '')} interrupted=True")
            yield {
                "type": "chat_completed",
                "content": partial_content,
                "reasoning": partial_reasoning or None,
                "agent_name": agent_name,
                "execution_order": execution_order,
                "assistant_message_id": 0,
                "assistant_msg_uuid": cancelled_msg_uuid,
                "interrupted": True
            }
        
        except Exception as e:
            logger.error(f"[AgentChatServiceV2] Error during chat: {e}", exc_info=True)
            yield {"type": "error", "message": str(e)}
        
        finally:
            # 清理 Todo 回调
            try:
                from .state.todo import TodoManager
                TodoManager.set_event_callback(None)
            except:
                pass
            
            if conversation_uuid in self._cancel_signals:
                del self._cancel_signals[conversation_uuid]
    
    def _get_tools_definition(self, agent_def=None) -> List[Dict[str, Any]]:
        """获取工具定义（包含 MCP 工具），并按 Agent 工具策略过滤

        agent_def.tools.enforce=True 时：被禁止工具的 schema 不会上传给 LLM；
        enforce=False 或 agent_def 为 None 时原样返回（向后兼容）。
        """
        try:
            tools = get_tools_definition_with_mcp(
                include_mcp=self._mcp_initialized,
                mcp_manager=self._mcp_manager
            )
        except Exception as e:
            logger.warning(f"[AgentChatServiceV2] Failed to get MCP tools, using base tools only: {e}")
            tools = get_tools_definition()
        # 按 Agent 权限过滤（enforce 未启用时原样返回）
        return filter_tools_for_agent(tools, agent_def)


# ==================== 单例 ====================

_instance: Optional[AgentChatServiceV2] = None


def get_agent_chat_service_v2() -> AgentChatServiceV2:
    """获取 AgentChatServiceV2 单例"""
    global _instance
    if _instance is None:
        _instance = AgentChatServiceV2()
    return _instance
